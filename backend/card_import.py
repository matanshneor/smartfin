"""ייבוא עסקאות מקובץ של חברת האשראי (מתן, 31.10).

פעם בחודש מורידים מהאתר של מקס, כאל או ישראכרט את פירוט העסקאות, ומעלים
אותו לכאן. המודול הזה טהור — בלי מסד ובלי רשת (חוץ מ-OpenAI, שמבודד בשתי
פונקציות בסוף) — כדי שאפשר יהיה לבדוק אותו עד הסוף בלי שום כפיל.

שלוש שכבות לזיהוי העמודות, מהזולה ליקרה:
  1. מילים מוכרות בשורת הכותרות ("תאריך עסקה", "שם בית העסק", "סכום חיוב").
     זה מכסה את כל החברות הגדולות בלי לדעת מראש מה המבנה המדויק שלהן.
  2. OpenAI — מקבל רק את שורת הכותרות ושלוש שורות לדוגמה, ואומר איזו עמודה
     היא מה (‎ai_map_columns‎).
  3. המשתמש בוחר בעצמו (‎mapping‎ שחוזר מהדפדפן).

מה שנשאר באחריות ‎app.py‎: הרשאות, כפילויות מול המסד, וההוספה עצמה.
"""
from __future__ import annotations

import calendar
import csv
import io
import json
import logging
import os
import re
from datetime import date, datetime, timedelta
from html.parser import HTMLParser

logger = logging.getLogger(__name__)

MAX_FILE_BYTES = 5 * 1024 * 1024
# שורות גולמיות שנסרקות בכל גיליון. פירוט חודשי של משפחה הוא עשרות עד
# מאות; התקרה היא נגד קובץ ענק שנשלח בטעות (או בכוונה).
MAX_ROWS_PER_SHEET = 5000
# עסקאות בייבוא אחד — שנה שלמה של משפחה פעילה נכנסת בזה בנוחות.
MAX_TRANSACTIONS = 500
# כמה ימים הפרש עדיין נחשבים "אותה עסקה" — הזנה ידנית נעשית לפעמים יום
# או יומיים אחרי הקנייה, וחברת האשראי רושמת לפעמים את יום הקליטה.
DUPLICATE_DAYS = 3

FIELDS = ("date", "merchant", "charge", "amount", "notes", "sector")


class ImportError_(Exception):
    """קובץ שאי אפשר לקרוא. ההודעה בעברית ומוצגת למשתמש."""


# ─── קריאת הקובץ לטבלאות ────────────────────────────────────────────────────

def read_sheets(data: bytes, filename: str = "") -> list:
    """הקובץ כרשימת גיליונות, כל גיליון — רשימת שורות, כל שורה — רשימת תאים.

    לפי התוכן ולא לפי הסיומת: חלק מהחברות מורידות "‎.xls‎" שהוא בפועל
    דף HTML עם טבלה, ואנשים משנים סיומות."""
    if not data:
        raise ImportError_("הקובץ ריק")
    if len(data) > MAX_FILE_BYTES:
        raise ImportError_("הקובץ גדול מדי (עד 5MB)")

    head = data[:8]
    if head.startswith(b"PK"):
        return _read_xlsx(data)
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        return _read_xls(data)

    if head.startswith(b"%PDF"):
        return _read_pdf(data)

    text = _decode(data)
    sniff = text.lstrip()[:2000].lower()
    if sniff.startswith("<") and ("<table" in sniff or "<html" in sniff or "<tr" in text.lower()):
        return [_read_html(text)]
    return [_read_csv(text)]


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1255"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_xlsx(data: bytes) -> list:
    try:
        import openpyxl
    except ImportError:                     # pragma: no cover — תלות מוצהרת
        raise ImportError_("קריאת קבצי אקסל אינה זמינה כרגע")
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        logger.info("card_import: xlsx לא נקרא", exc_info=True)
        raise ImportError_("לא הצלחנו לפתוח את קובץ האקסל — ודאו שזה הקובץ שהורד מהאתר")
    sheets = []
    try:
        for ws in wb.worksheets:
            rows = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_ROWS_PER_SHEET:
                    break
                rows.append([_clean_cell(c) for c in row])
            sheets.append(rows)
    finally:
        wb.close()
    return sheets


def _read_xls(data: bytes) -> list:
    try:
        import xlrd
    except ImportError:                     # pragma: no cover — תלות מוצהרת
        raise ImportError_("קריאת קבצי אקסל ישנים אינה זמינה כרגע")
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        logger.info("card_import: xls לא נקרא", exc_info=True)
        raise ImportError_("לא הצלחנו לפתוח את קובץ האקסל — ודאו שזה הקובץ שהורד מהאתר")
    sheets = []
    for sh in book.sheets():
        rows = []
        for r in range(min(sh.nrows, MAX_ROWS_PER_SHEET)):
            cells = []
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
                        continue
                    except Exception:
                        pass
                cells.append(_clean_cell(cell.value))
            rows.append(cells)
        sheets.append(rows)
    return sheets


class _TableParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.rows, self._row, self._cell = [], None, None

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th") and self._row is not None:
            self._cell = []
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self._row is not None and self._cell is not None:
            self._row.append(_clean_cell("".join(self._cell)))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if len(self.rows) < MAX_ROWS_PER_SHEET:
                self.rows.append(self._row)
            self._row = None

    def handle_data(self, data):
        if self._cell is not None:
            self._cell.append(data)


def _read_html(text: str) -> list:
    parser = _TableParser()
    try:
        parser.feed(text)
        parser.close()
    except Exception:
        raise ImportError_("לא הצלחנו לקרוא את הקובץ")
    return parser.rows


def _read_csv(text: str) -> list:
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= MAX_ROWS_PER_SHEET:
            break
        rows.append([_clean_cell(c) for c in row])
    return rows


# ─── PDF: דף החיוב החודשי (מתן, 31.10 — דוגמה של כאל) ───────────────────────
#
# ב-PDF אין טבלה — יש מילים במקומות על הדף. משחזרים אותה כך:
#   1. מילים → שורות, לפי הגובה בדף.
#   2. "שורת נתונים" = שורה עם תאריך וסכום. רצף שלהן הוא טבלה.
#   3. הכותרות (שורה אחת או שתיים צמודות מעל — "תאריך / העסקה") קובעות את
#      העמודות לפי המיקום האופקי; כל מילה בשורת נתונים הולכת לעמודה שלה.
# כל טבלה יוצאת כ"גיליון" נפרד, ומשם הכל כמו באקסל — אותו זיהוי כותרות.
#
# ובעברית: חלק מהקבצים שומרים כל מילה בסדר חזותי — "העסק" נשמר "קסעה".
# מזהים את זה לפי מילות הכותרת המוכרות, והופכים.

MAX_PDF_PAGES = 30
_HEB = re.compile(r"[֐-׿]")
_HEB_RUNS = re.compile(r"[֐-׿\"'׳״]+|[^֐-׿\"'׳״]+")
_KNOWN_WORDS = ("תאריך", "סכום", "העסק", "העסקה", "חיוב", "פירוט", "עסקות", "עסקאות")
# כמה נקודות מתחת לשורה קודמת עוד נחשבות אותה שורה, ומרווחים בין שורות כותרת
_LINE_TOLERANCE = 2.0
_HEADER_TO_DATA = 15.0
_HEADER_TO_HEADER = 8.0
_BLOCK_GAP = 40.0
_COLUMN_GAP = 6.0


def _read_pdf(data: bytes) -> list:
    try:
        import pdfplumber
    except ImportError:                     # pragma: no cover — תלות מוצהרת
        raise ImportError_("קריאת PDF אינה זמינה כרגע — הורידו מהאתר את הפירוט כקובץ אקסל")
    pages = []
    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages[:MAX_PDF_PAGES]:
                pages.append([{"text": w["text"], "x0": w["x0"], "x1": w["x1"], "top": w["top"]}
                              for w in page.extract_words()])
    except Exception:
        logger.info("card_import: pdf לא נקרא", exc_info=True)
        raise ImportError_("לא הצלחנו לפתוח את קובץ ה-PDF — נסו להוריד מהאתר את הפירוט כקובץ אקסל")
    sheets = sheets_from_pdf_words(pages)
    if not sheets:
        raise ImportError_("לא מצאנו טקסט ב-PDF (אולי זה קובץ סרוק). "
                           "הורידו מהאתר את הפירוט כקובץ אקסל, או כ-PDF מקורי")
    return sheets


def _visual_order(pages) -> bool:
    """האם המילים העבריות שמורות הפוך. מכריעים לפי מילות כותרת מוכרות."""
    straight = flipped = 0
    for words in pages:
        for w in words:
            t = w["text"]
            if not _HEB.search(t):
                continue
            straight += any(k in t for k in _KNOWN_WORDS)
            flipped += any(k in t[::-1] for k in _KNOWN_WORDS)
    return flipped > straight


def _fix_word(text: str, visual: bool) -> str:
    if not visual or not _HEB.search(text):
        return text
    runs = _HEB_RUNS.findall(text)
    return "".join(r[::-1] if _HEB.search(r) else r for r in reversed(runs))


def _lines(words, visual):
    """מילים → שורות, מלמעלה למטה; בכל שורה המילים מימין לשמאל."""
    lines = []
    for w in sorted(words, key=lambda w: w["top"]):
        w = dict(w, text=_fix_word(w["text"], visual))
        if lines and abs(lines[-1]["top"] - w["top"]) <= _LINE_TOLERANCE:
            lines[-1]["words"].append(w)
        else:
            lines.append({"top": w["top"], "words": [w]})
    for line in lines:
        line["words"].sort(key=lambda w: -w["x1"])
        line["text"] = " ".join(w["text"] for w in line["words"])
    return lines


def _is_data_line(line) -> bool:
    texts = [w["text"] for w in line["words"]]
    has_date = any(parse_date(t) for t in texts)
    has_amount = any(parse_date(t) is None and _parse_amount_or_none(t) not in (None, 0)
                     and re.search(r"\d[.,]\d{2}$|^\d+$", t.strip("₪ ")) for t in texts)
    return has_date and has_amount


def _columns(header_lines):
    """עמודות מהכותרות: מילים צמודות באותה שורה הן כותרת אחת ("שם בית
    העסק"), וכותרות בשתי שורות שחופפות אופקית הן אותה עמודה ("סכום / חיוב")."""
    cols = []
    for line in header_lines:
        clusters = []
        for w in line["words"]:              # מימין לשמאל
            if clusters and clusters[-1]["x0"] - w["x1"] <= _COLUMN_GAP:
                c = clusters[-1]
                c["x0"] = min(c["x0"], w["x0"])
                c["parts"].append(w["text"])
            else:
                clusters.append({"x0": w["x0"], "x1": w["x1"], "parts": [w["text"]]})
        for c in clusters:
            same = next((k for k in cols if c["x0"] < k["x1"] and k["x0"] < c["x1"]), None)
            if same:
                same["x0"], same["x1"] = min(same["x0"], c["x0"]), max(same["x1"], c["x1"])
                same["text"] += " " + " ".join(c["parts"])
            else:
                cols.append({"x0": c["x0"], "x1": c["x1"], "text": " ".join(c["parts"])})
    return sorted(cols, key=lambda c: -c["x1"])


def _cells(line, cols):
    """כל מילה לעמודה הקרובה אליה (לפי הגבול שבאמצע הרווח בין עמודות)."""
    cells = [[] for _ in cols]
    for w in line["words"]:
        mid = (w["x0"] + w["x1"]) / 2
        best = min(range(len(cols)), key=lambda i: 0 if cols[i]["x0"] <= mid <= cols[i]["x1"]
                   else min(abs(mid - cols[i]["x0"]), abs(mid - cols[i]["x1"])))
        cells[best].append(w["text"])
    return [" ".join(c) if c else None for c in cells]


def sheets_from_pdf_words(pages) -> list:
    """עמודים של מילים (‎text, x0, x1, top‎) → "גיליונות" לפי טבלאות.
    נפרד מ-‎pdfplumber‎ כדי שאפשר יהיה לבדוק אותו בלי קובץ."""
    visual = _visual_order(pages)
    sheets = []
    for words in pages:
        lines = _lines(words, visual)
        i = 0
        while i < len(lines):
            if not _is_data_line(lines[i]):
                i += 1
                continue
            # הכותרות: עד שתי שורות צמודות מעל שורת הנתונים הראשונה
            header, top = [], lines[i]["top"]
            j = i - 1
            gap = _HEADER_TO_DATA
            while j >= 0 and len(header) < 2 and top - lines[j]["top"] <= gap \
                    and not _is_data_line(lines[j]):
                header.insert(0, lines[j])
                top, gap = lines[j]["top"], _HEADER_TO_HEADER
                j -= 1
            block, last = [], lines[i]["top"]
            while i < len(lines) and lines[i]["top"] - last <= _BLOCK_GAP:
                if _is_data_line(lines[i]):
                    block.append(lines[i])
                    last = lines[i]["top"]
                i += 1
            if not header:
                continue
            cols = _columns(header)
            sheets.append([[c["text"] for c in cols]] + [_cells(l, cols) for l in block])
    return sheets


# סימני כיווניות שאקסל בעברית מפזר בתוך תאים — בלתי נראים, ושוברים השוואות
_BIDI = dict.fromkeys(map(ord, "‎‏‪‫‬‭‮⁦⁧⁨⁩﻿"), None)


def _clean_cell(value):
    if value is None:
        return None
    if isinstance(value, str):
        value = value.translate(_BIDI).replace("\xa0", " ").strip()
        return value or None
    return value


# ─── זיהוי העמודות ──────────────────────────────────────────────────────────

def _norm_header(value) -> str:
    if not isinstance(value, str):
        return ""
    v = value.translate(_BIDI).lower()
    v = re.sub(r"[\"'״׳`]", "", v)
    v = re.sub(r"[\s_:.\-()]+", " ", v)
    return v.strip()


def _first(headers, used, *preds):
    """העמודה הראשונה שמתאימה לתנאי הראשון שמתאים למשהו — כך ש"תאריך
    עסקה" מנצח את "תאריך חיוב" גם כשהוא מופיע אחריו."""
    for pred in preds:
        for i, h in enumerate(headers):
            if h and i not in used and pred(h):
                return i
    return None


def detect_mapping(header_row) -> dict | None:
    """מיפוי שדה → אינדקס עמודה לפי המילים בכותרות. ‎None‎ — לא שורת כותרות.

    חייבים תאריך, בית עסק וסכום כלשהו. "סכום חיוב" (בשקלים, מה שבאמת ירד)
    עדיף על "סכום עסקה" (שבעסקת חו"ל הוא בדולרים)."""
    headers = [_norm_header(c) for c in header_row]
    if sum(1 for h in headers if h) < 3:
        return None
    used = set()

    def take(i):
        if i is not None:
            used.add(i)
        return i

    has = lambda *words: (lambda h: all(w in h for w in words))
    not_ = lambda pred, *words: (lambda h: pred(h) and not any(w in h for w in words))

    m = {}
    m["date"] = take(_first(headers, used,
                            lambda h: h in ("תאריך עסקה", "תאריך העסקה", "תאריך רכישה",
                                            "תאריך ביצוע", "תאריך", "date", "transaction date"),
                            not_(has("תאריך"), "חיוב", "קליטה"),
                            has("date"),
                            has("תאריך")))
    m["charge"] = take(_first(headers, used,
                              not_(has("סכום", "חיוב"), "מקורי"),
                              has("חיוב", "שח"),
                              has("charged"),
                              lambda h: h in ("חיוב", "סכום בשח", "סכום בש ח")))
    m["amount"] = take(_first(headers, used,
                              lambda h: h in ("סכום עסקה", "סכום העסקה", "סכום עסקה מקורי",
                                              "סכום מקורי", "סכום", "amount"),
                              not_(has("סכום"), "מטבע", "שער"),
                              has("amount")))
    m["merchant"] = take(_first(headers, used,
                                has("בית", "עסק"),
                                has("שם העסק"),
                                lambda h: h in ("תיאור", "תאור", "פרטים", "שם", "description",
                                                "merchant", "תיאור התנועה", "פרטי העסקה"),
                                not_(has("תיאור"), "נוסף"),
                                has("merchant"), has("description")))
    m["notes"] = take(_first(headers, used,
                             lambda h: h in ("הערות", "הערה", "פירוט נוסף", "סוג עסקה",
                                             "פרטים נוספים", "notes"),
                             has("הער"), has("סוג עסק"), has("פירוט")))
    # הקטגוריה של חברת האשראי עצמה ("ענף" בכאל, "קטגוריה" במקס) — רמז
    m["sector"] = take(_first(headers, used,
                              lambda h: h in ("ענף", "קטגוריה", "ענף עסקי", "סוג בית עסק",
                                              "category")))
    if m["date"] is None or m["merchant"] is None or (m["charge"] is None and m["amount"] is None):
        return None
    return {k: v for k, v in m.items() if v is not None}


def header_signature(row) -> list:
    return [_norm_header(c) for c in row]


def find_header_candidate(sheets):
    """לקובץ שהמילים לא זיהו: השורה שהכי סביר שהיא הכותרות, ושלוש שורות
    אחריה — מה שנשלח ל-AI ומה שמוצג למשתמש בבחירה הידנית.

    מוצאים את שורת הנתונים הראשונה (יש בה תאריך וסכום), והכותרת היא השורה
    הקרובה מעליה שיש בה לפחות שני תאי טקסט."""
    for rows in sheets:
        for d, data_row in enumerate(rows[:80]):
            if not _row_has_date_and_amount(data_row):
                continue
            for i in range(d - 1, max(d - 6, -1), -1):
                row = rows[i]
                texts = [c for c in row if isinstance(c, str) and parse_date(c) is None
                         and _parse_amount_or_none(c) is None]
                if len(texts) < 2:
                    continue
                following = [r for r in rows[i + 1:i + 8] if any(c is not None for c in r)][:3]
                width = max([len(row)] + [len(r) for r in following])
                return {
                    "signature": header_signature(row),
                    "headers": [_cell_text(c) for c in row] + [""] * (width - len(row)),
                    "samples": [[_cell_text(c) for c in r] + [""] * (width - len(r))
                                for r in following],
                }
            break
    return None


def _row_has_date_and_amount(row) -> bool:
    return any(parse_date(c) for c in row) and any(
        _parse_amount_or_none(c) not in (None, 0) and parse_date(c) is None for c in row)


def _cell_text(c) -> str:
    if c is None:
        return ""
    if isinstance(c, datetime):
        return c.strftime("%d/%m/%Y")
    if isinstance(c, date):
        return c.strftime("%d/%m/%Y")
    if isinstance(c, float) and c.is_integer():
        return str(int(c))
    return str(c)[:80]


def valid_mapping(mapping, width: int) -> dict | None:
    """מיפוי שהגיע מבחוץ (AI או הדפדפן): אינדקסים שלמים בתוך רוחב הטבלה,
    בלי עמודה כפולה, ועם השדות החובה."""
    if not isinstance(mapping, dict):
        return None
    out = {}
    for k in FIELDS:
        v = mapping.get(k)
        if v is None or v == "" or v == -1:
            continue
        if isinstance(v, bool) or not isinstance(v, int) or not 0 <= v < max(width, 1):
            return None
        out[k] = v
    if len(set(out.values())) != len(out):
        return None
    if "date" not in out or "merchant" not in out or not ({"charge", "amount"} & out.keys()):
        return None
    return out


# ─── ערכים ──────────────────────────────────────────────────────────────────

_EXCEL_EPOCH = date(1899, 12, 30)
_DMY = re.compile(r"^(\d{1,2})[./\-](\d{1,2})[./\-](\d{2}|\d{4})(?:\s.*)?$")
_YMD = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[T\s].*)?$")


def parse_date(value):
    """תאריך מתא. סדר ישראלי (יום/חודש/שנה) — אף חברה ישראלית לא מוציאה
    חודש/יום, ונחש לא נכון היה מזיז עסקאות בין חודשים בשקט."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        # מספר סידורי של אקסל, רק בטווח הגיוני (1954–2119)
        if 20000 <= value <= 80000:
            return _EXCEL_EPOCH + timedelta(days=int(value))
        return None
    if not isinstance(value, str):
        return None
    s = value.strip()
    m = _DMY.match(s)
    try:
        if m:
            d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
            if y < 100:
                y += 2000
            return date(y, mo, d)
        m = _YMD.match(s)
        if m:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    except ValueError:
        return None
    return None


_NUM = re.compile(r"-?\d+(?:\.\d+)?")


def _parse_amount_or_none(value):
    """סכום מתא, עם סימן. "₪ 1,234.50", "54.90-", "(54.90)" ו-"‎-54.9‎"."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if not isinstance(value, str):
        return None
    s = value.translate(_BIDI).strip()
    if not s or not re.search(r"\d", s):
        return None
    # תא עם אותיות שאינן סימון מטבע — טקסט, לא סכום ("תשלום 3 מתוך 12")
    letters = re.sub(r"(ש\"?ח|ש״ח|nis|ils|usd|eur|₪|\$|€|£)", "", s, flags=re.I)
    if re.search(r"[A-Za-z֐-׿]", letters):
        return None
    if _DMY.match(s) or _YMD.match(s):
        return None
    negative = s.startswith("-") or s.endswith("-") or (s.startswith("(") and s.endswith(")"))
    digits = re.sub(r"[^\d.\-]", "", s.replace(",", ""))
    found = _NUM.search(digits.replace("-", ""))
    if not found:
        return None
    try:
        number = float(found.group(0))
    except ValueError:
        return None
    return -number if negative else number


parse_amount = _parse_amount_or_none

_INSTALLMENT = re.compile(r"תשלום\s*(\d{1,2})\s*(?:מתוך|מ-|מ\s|/|מ)\s*(\d{1,2})")


def parse_installment(cells):
    """"תשלום 3 מתוך 12" באחד התאים — ‎(3, 12)‎, או ‎None‎."""
    for c in cells:
        if isinstance(c, str):
            m = _INSTALLMENT.search(c)
            if m:
                k, n = int(m.group(1)), int(m.group(2))
                if 1 <= k <= n <= 60 and n >= 2:
                    return k, n
    return None


def _add_months(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    y, m = d.year + y, m + 1
    return date(y, m, min(d.day, calendar.monthrange(y, m)[1]))


# ─── חילוץ העסקאות ──────────────────────────────────────────────────────────

_SUMMARY_WORDS = ("סהכ", "סה כ", "סך הכל", "total", "יתרה")


def extract(sheets, known_layouts=None):
    """כל העסקאות בקובץ.

    ‎known_layouts‎: מיפויים שכבר אושרו (מה-AI או מהמשתמש), כרשימה של
    ‎{"signature": [...], "mapping": {...}}‎. שורה שהכותרות שלה זהות לחתימה
    מקבלת את המיפוי הזה.

    מחזיר ‎(transactions, layouts_used)‎. רשימה ריקה של עסקאות עם
    ‎layouts_used‎ ריק — לא נמצאה שורת כותרות; ‎app.py‎ עובר לשכבה הבאה."""
    known = []
    for layout in known_layouts or []:
        if isinstance(layout, dict) and isinstance(layout.get("signature"), list):
            known.append(([str(s) for s in layout["signature"]], layout.get("mapping")))

    out, used = [], []
    for rows in sheets:
        mapping = None
        for row in rows:
            if not row or all(c is None for c in row):
                continue
            sig = header_signature(row)
            new_mapping = None
            for ksig, kmap in known:
                if _same_signature(sig, ksig):
                    new_mapping = valid_mapping(kmap, len(row))
                    break
            if new_mapping is None:
                new_mapping = detect_mapping(row)
            if new_mapping is not None:
                mapping = new_mapping
                entry = {"signature": sig, "mapping": mapping}
                if entry not in used:
                    used.append(entry)
                continue
            if mapping is None:
                continue
            tx = _row_to_tx(row, mapping)
            if tx:
                out.append(tx)
                if len(out) > MAX_TRANSACTIONS:
                    raise ImportError_(f"יש בקובץ יותר מ-{MAX_TRANSACTIONS} עסקאות — "
                                       f"הורידו פירוט של תקופה קצרה יותר")
    return out, used


def _same_signature(a, b) -> bool:
    a = [x for x in a]
    b = [x for x in b]
    while a and not a[-1]:
        a.pop()
    while b and not b[-1]:
        b.pop()
    return a == b and any(a)


def _cell(row, i):
    return row[i] if i is not None and i < len(row) else None


def _row_to_tx(row, mapping):
    when = parse_date(_cell(row, mapping.get("date")))
    if when is None:
        return None
    merchant = _cell(row, mapping.get("merchant"))
    if isinstance(merchant, (int, float)) and not isinstance(merchant, bool):
        merchant = _cell_text(merchant)
    if not isinstance(merchant, str) or not merchant.strip():
        return None
    merchant = re.sub(r"\s+", " ", merchant).strip()[:100]
    if any(w in _norm_header(merchant) for w in _SUMMARY_WORDS):
        return None

    charge = parse_amount(_cell(row, mapping.get("charge")))
    amount = parse_amount(_cell(row, mapping.get("amount")))
    # עסקה שעוד לא חויבה (או בחו"ל, לפני המרה) — "סכום חיוב" ריק או 0
    value = charge if charge not in (None, 0) else amount
    if value in (None, 0):
        return None

    inst = parse_installment([_cell(row, mapping.get("notes"))] + list(row))
    description = merchant
    if inst:
        k, n = inst
        description = f"{merchant} (תשלום {k} מתוך {n})"
        # בתשלומים, "תאריך עסקה" הוא יום הקנייה בכל אחד מהתשלומים. התשלום
        # ה-k יורד k-1 חודשים אחריו — אחרת כל 12 התשלומים היו נוחתים באותו
        # חודש, ובייבוא של החודש הבא — כפילות שלא נתפסת.
        when = _add_months(when, k - 1)

    return {
        "date":        when.isoformat(),
        "merchant":    merchant,
        "description": description[:120],
        "amount":      round(abs(value), 2),
        "refund":      value < 0,
        "installment": list(inst) if inst else None,
        "sector":      _sector(_cell(row, mapping.get("sector"))),
    }


def _sector(value):
    return re.sub(r"\s+", " ", value).strip()[:40] if isinstance(value, str) and value.strip() else None


# ─── שמות בתי עסק: התאמה להיסטוריה ──────────────────────────────────────────

def merchant_key(text) -> str:
    """שם בית עסק להשוואה: בלי ספרות (מספר סניף), סימנים ורווחים כפולים.
    "פז חולון 4531" ו"פז  חולון" — אותו מפתח."""
    if not isinstance(text, str):
        return ""
    t = text.translate(_BIDI).lower()
    t = re.sub(r"\(תשלום[^)]*\)", " ", t)
    t = re.sub(r"[\d]+", " ", t)
    t = re.sub(r"[^\w\s]", " ", t)
    t = re.sub(r"_", " ", t)
    return re.sub(r"\s+", " ", t).strip()


def _related(a: str, b: str) -> bool:
    """אותו בית עסק: מפתח זהה, או שאחד הוא תחילת השני במילים שלמות
    ("שופרסל" שהוזן ידנית ו"שופרסל דיל" מהקובץ)."""
    if not a or not b:
        return False
    if a == b:
        return True
    short, long_ = (a, b) if len(a) <= len(b) else (b, a)
    return len(short) >= 3 and long_.startswith(short + " ")


def learned_categories(history) -> dict:
    """מפתח בית עסק → הקטגוריה האחרונה שנבחרה לו במשפחה.

    ‎history‎: שורות ‎{description, category_id, date}‎, מהחדשה לישנה. זה כל
    ה"זיכרון": עסקה שמיובאת עם קטגוריה שתוקנה נשמרת עם תיאור = שם בית העסק,
    ובחודש הבא היא זו שנמצאת כאן. גם עסקאות שהוזנו ידנית מלמדות."""
    out = {}
    for row in history:
        key = merchant_key(row.get("description"))
        if key and row.get("category_id") and key not in out:
            out[key] = row["category_id"]
    return out


def lookup_learned(key: str, learned: dict):
    if key in learned:
        return learned[key]
    best = None
    for k, cat in learned.items():
        if _related(key, k) and (best is None or len(k) > len(best[0])):
            best = (k, cat)
    return best[1] if best else None


# כללים לפעם הראשונה, לפני שיש היסטוריה: מילה בשם בית העסק → מילה בשם
# הקטגוריה. לפי שם ולא לפי מזהה, כי לכל משפחה קטגוריות משלה — וכלל חל רק
# כשיש למשפחה קטגוריה שהשם שלה מתאים ("סופר ומזון", "מזון").
_RULES = [
    (("שופרסל", "רמי לוי", "ויקטורי", "יוחננוף", "אושר עד", "טיב טעם", "יינות ביתן",
      "חצי חינם", "קרפור", "carrefour", "מגה", "פרש מרקט", "סופר", "מכולת", "מאפי",
      "קצביה", "ירקות", "am pm", "ampm", "שוק"),
     ("מזון", "סופר")),
    (("פז", "דלק", "סונול", "דור אלון", "ten", "yellow", "רב קו", "gett", "moovit",
      "מוביט", "חניון", "חניה", "פנגו", "pango", "cellopark", "סלופארק", "כביש",
      "מוסך", "רכבת", "אגד", "דן"),
     ("רכב", "תחבורה")),
    (("חברת החשמל", "חשמל", "מים", "תאגיד", "ארנונה", "עירית", "עיריית", "בזק", "hot",
      "הוט", "yes", "פרטנר", "partner", "סלקום", "cellcom", "פלאפון", "pelephone",
      "גולן טלקום", "netflix", "spotify", "apple com", "google", "דיסני", "disney", "ביטוח"),
     ("חשבונות", "מנויים")),
    (("סופר פארם", "super pharm", "מכבי", "כללית", "מאוחדת", "לאומית", "בית מרקחת",
      "פארם", "רופא", "שיניים", "אופטיקה"),
     ("בריאות",)),
    (("מסעד", "קפה", "cafe", "coffee", "ארומה", "aroma", "wolt", "וולט", "תן ביס", "10bis",
      "cibus", "סיבוס", "מקדונלד", "mcdonald", "בורגר", "burger", "פיצ", "pizza", "דומינו",
      "סינמה", "cinema", "יס פלאנט", "פאב", "בר"),
     ("מסעדות", "בילויים", "בידור")),
    (("zara", "זארה", "h m", "קסטרו", "castro", "fox", "פוקס", "aliexpress", "amazon",
      "אמזון", "shein", "שיין", "terminal x", "טרמינל", "אדידס", "adidas", "נייקי", "nike",
      "ksp", "באג", "מחסני חשמל"),
     ("שופינג", "קניות", "ביגוד")),
    (("ikea", "איקאה", "ace", "אייס", "הום סנטר", "home center", "ביתילי", "פוקס הום"),
     ("ריהוט", "לבית", "בית")),
    (("אל על", "el al", "booking", "airbnb", "ישראייר", "ארקיע", "hotel", "מלון", "wizz",
      "ryanair", "easyjet"),
     ("חופשות", "טיולים", "נסיעות")),
    (("מספרה", "קוסמטיק", "ציפורניים", "nails", "spa", "ספא"),
     ("טיפוח",)),
]


def rule_category(merchant: str, categories: list):
    """קטגוריה לפי הכללים, מתוך קטגוריות ההוצאה של המשפחה. ‎None‎ — אין.

    מילה קצרה (עד 3 תווים: "פז", "בר", "hot") חייבת להופיע כמילה שלמה —
    אחרת "בר" היה תופס את "עבר" ו"פז" את "חפז"."""
    text = (merchant or "").lower().translate(_BIDI)
    text = " " + re.sub(r"\s+", " ", re.sub(r"[^\w\s]", " ", text)).strip() + " "
    # המילה הארוכה ביותר מנצחת: "סופר פארם" הוא בריאות ולא "סופר",
    # ו"מחסני חשמל" הוא קניות ולא "חשמל"
    best = None
    for merchant_words, category_words in _RULES:
        hits = [w for w in merchant_words if ((f" {w} " in text) if len(w) <= 3 else (w in text))]
        if not hits:
            continue
        cat = next((c for c in categories
                    if any(cw in (c.get("name") or "") for cw in category_words)), None)
        if cat and (best is None or len(max(hits, key=len)) > best[0]):
            best = (len(max(hits, key=len)), cat["id"])
    return best[1] if best else None


# הענף שחברת האשראי עצמה נותנת ("פנאי בילוי", "רכב ותחבורה") → מילה בשם
# הקטגוריה של המשפחה. רק כשאין התאמה לפי בית העסק עצמו — הענף כללי יותר.
_SECTOR_RULES = [
    (("מזון", "סופר", "מכולת"), ("מזון", "סופר")),
    (("רכב", "תחבור", "דלק", "חניה"), ("רכב", "תחבורה")),
    (("מסעד", "בילוי", "פנאי", "מזון מהיר", "בתי קפה"), ("מסעדות", "בילויים", "בידור")),
    (("תקשורת", "מחשבים", "חשמל ומים", "ביטוח", "שירותים"), ("חשבונות", "מנויים")),
    (("רפואה", "בריאות", "פארם", "אופטיקה"), ("בריאות",)),
    (("אופנה", "ביגוד", "הנעלה", "קניות", "כלבו", "מוצרי און"), ("שופינג", "קניות", "ביגוד")),
    (("ריהוט", "בית ומשרד", "בית וגן", "עיצוב"), ("ריהוט", "לבית", "בית")),
    (("תיירות", "טיסות", "נופש", "מלונות"), ("חופשות", "טיולים", "נסיעות")),
    (("קוסמטיקה", "טיפוח", "יופי"), ("טיפוח",)),
]


def sector_category(sector, categories: list):
    if not sector:
        return None
    for sector_words, category_words in _SECTOR_RULES:
        if any(w in sector for w in sector_words):
            cat = next((c for c in categories
                        if any(cw in (c.get("name") or "") for cw in category_words)), None)
            if cat:
                return cat["id"]
    return None


# ─── כפילויות ───────────────────────────────────────────────────────────────

def mark_duplicates(items: list, existing: list) -> None:
    """מסמן כל עסקה מהקובץ שכבר קיימת במסד. משנה את ‎items‎ במקום.

    ‎exists‎: אותו סכום, עד ‎DUPLICATE_DAYS‎ ימים הפרש, ואותו בית עסק (או
    שבעסקה הקיימת אין תיאור). ‎maybe‎: אותו סכום ותאריך קרוב, אבל תיאור
    אחר — אולי הוזנה ידנית בשם אחר; לא מסומנת לייבוא, והמשתמש מחליט.

    כל שורה קיימת "נצרכת" פעם אחת: שני קפה זהים באותו יום בקובץ, ואחד
    מהם הוזן ידנית — השני עדיין חדש."""
    pool = []
    for row in existing:
        try:
            d = date.fromisoformat(str(row.get("date"))[:10])
            amt = round(float(row.get("amount")), 2)
        except (TypeError, ValueError):
            continue
        pool.append({"date": d, "amount": amt, "key": merchant_key(row.get("description")),
                     "description": row.get("description") or "", "used": False})

    for want_exact in (True, False):
        for item in items:
            if item.get("status"):
                continue
            d = date.fromisoformat(item["date"])
            key = merchant_key(item["merchant"])
            best = None
            for p in pool:
                if p["used"] or abs(p["amount"] - item["amount"]) > 0.005:
                    continue
                gap = abs((p["date"] - d).days)
                if gap > DUPLICATE_DAYS:
                    continue
                same = not p["key"] or _related(key, p["key"])
                if want_exact != same:
                    continue
                if best is None or gap < best[0]:
                    best = (gap, p)
            if best:
                p = best[1]
                p["used"] = True
                item["status"] = "exists" if want_exact else "maybe"
                item["match"] = {"date": p["date"].isoformat(), "description": p["description"]}
    for item in items:
        item.setdefault("status", "new")


# ─── OpenAI ─────────────────────────────────────────────────────────────────

def _openai():
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        return None
    try:
        from openai import OpenAI
    except ImportError:                     # pragma: no cover
        return None
    return OpenAI(api_key=api_key, timeout=25)


def _call_tool(client, tool, prompt, max_tokens):
    response = client.chat.completions.create(
        model="gpt-4o-mini",
        max_tokens=max_tokens,
        tools=[tool],
        tool_choice={"type": "function", "function": {"name": tool["function"]["name"]}},
        messages=[{"role": "user", "content": prompt}],
    )
    calls = response.choices[0].message.tool_calls
    if not calls:
        return None
    try:
        return json.loads(calls[0].function.arguments or "{}")
    except (TypeError, ValueError):
        return None


def ai_map_columns(headers: list, samples: list):
    """שכבה 2: איזו עמודה היא מה, לפי הכותרות ושלוש שורות לדוגמה בלבד.
    לא נשלח שום דבר מעבר לזה. ‎None‎ — לא זמין או לא הצליח."""
    client = _openai()
    if not client:
        return None
    col = {"type": "integer", "description": "אינדקס העמודה (מ-0), או ‎-1‎ אם אין כזאת"}
    tool = {"type": "function", "function": {
        "name": "map_columns",
        "description": "ממפה עמודות של קובץ פירוט עסקאות אשראי ישראלי לשדות.",
        "parameters": {"type": "object", "properties": {
            "date": dict(col, description="תאריך העסקה (לא תאריך החיוב), אינדקס מ-0"),
            "merchant": dict(col, description="שם בית העסק / תיאור העסקה, אינדקס מ-0"),
            "charge": dict(col, description="הסכום שחויב בשקלים, אינדקס מ-0, או -1"),
            "amount": dict(col, description="סכום העסקה המקורי, אינדקס מ-0, או -1"),
            "notes": dict(col, description="הערות / סוג עסקה (למשל תשלומים), או -1"),
        }, "required": ["date", "merchant", "charge", "amount"]},
    }}
    table = "\n".join(" | ".join(f"[{i}] {h}" for i, h in enumerate(r))
                      for r in [headers] + samples)
    try:
        data = _call_tool(client, tool,
                          "אלה הכותרות ושלוש שורות ראשונות מקובץ פירוט עסקאות של כרטיס אשראי. "
                          "החזר באמצעות הכלי איזו עמודה היא כל שדה.\n\n" + table, 200)
    except Exception:
        logger.exception("ai_map_columns")
        return None
    if not isinstance(data, dict):
        return None
    return valid_mapping({k: data.get(k) for k in FIELDS}, len(headers))


def ai_categorize(merchants: list, category_names: list) -> dict:
    """שמות בתי עסק → שם קטגוריה מתוך הקטגוריות של המשפחה בלבד.
    רק שמות בתי העסק נשלחים — בלי סכומים, תאריכים או פרטי כרטיס."""
    if not merchants or not category_names:
        return {}
    client = _openai()
    if not client:
        return {}
    merchants = merchants[:80]
    tool = {"type": "function", "function": {
        "name": "categorize",
        "description": "משייך כל בית עסק ישראלי לקטגוריית הוצאה.",
        "parameters": {"type": "object", "properties": {
            "items": {"type": "array", "items": {"type": "object", "properties": {
                "merchant": {"type": "string"},
                "category": {"type": "string", "enum": category_names + [""],
                             "description": "ריק אם אין התאמה ברורה"},
            }, "required": ["merchant", "category"]}},
        }, "required": ["items"]},
    }}
    try:
        data = _call_tool(client, tool,
                          "שייך כל אחד מבתי העסק הבאים (מפירוט כרטיס אשראי בישראל) לאחת "
                          "הקטגוריות. אם לא ברור — השאר ריק.\n\n" + "\n".join(merchants),
                          60 + 40 * len(merchants))
    except Exception:
        logger.exception("ai_categorize")
        return {}
    out = {}
    allowed = set(category_names)
    for item in (data or {}).get("items") or []:
        if isinstance(item, dict) and item.get("category") in allowed \
                and isinstance(item.get("merchant"), str):
            out[item["merchant"]] = item["category"]
    return out
